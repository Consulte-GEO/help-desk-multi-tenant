# -*- coding: utf-8 -*-
"""Preenche uma CÓPIA do R02 com dados de exemplo (sintéticos) para teste
da importação. Preserva estilos/dropdowns do modelo — só escreve valores."""
import shutil
from openpyxl import load_workbook

SRC = 'frontend/public/modelos/Modelo-Implantacao-Consulte-GEO-R02.xlsx'
DST = 'frontend/public/modelos/Modelo-Implantacao-Consulte-GEO-R02-EXEMPLO-PREENCHIDO.xlsx'
shutil.copy(SRC, DST)
wb = load_workbook(DST)

obra = [
    ('Empresa / Contratante *', 'Construtora Exemplo Ltda'),
    ('CNPJ', '12.345.678/0001-90'),
    ('Nome da Obra *', 'Residencial Vista Verde'),
    ('Endereço completo *', 'Rua das Acácias, 1200 - Jardim Europa'),
    ('CEP', '01234-567'),
    ('Cidade / Regional', 'São Paulo - SP'),
    ('Quantidade de blocos/torres', '2 torres + área externa'),
    ('Número de unidades', 168),
    ('Engenheiro da obra', 'Ana Paula Souza'),
]
ws = wb['Obra']
for r in range(3, 3 + len(obra)):
    ws.cell(row=r, column=2, value=obra[r - 3][1])

ws = wb['Estruturas']
for i, row in enumerate([
    ('Torre', '1 a 2', 'Torres 1 e 2', 'Projetos Estruturais Alfa', ''),
    ('Área Externa', '-', 'Área Externa', '', 'muros, passeio e guarita'),
    ('Contenção', '-', 'Contenções', 'Geotecnia Beta', 'talude fundos'),
], start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=i, column=c, value=v)

ws = wb['Concretos']
for i, row in enumerate([
    ('FCK 25 | Bombeável - Slump 120±20mm - Brita 1', 'Fundação', 25, '1 CP - 7 dias / 2 CPs - 28 dias'),
    ('FCK 30 | Bombeável - Slump 120±20mm - Brita 0+1', 'Lajes, Vigas e Pilares', 30, '1 CP - 7 dias / 2 CPs - 28 dias / 1 CP - 63 dias'),
    ('FCK 40 | Bombeável - Slump 120±20mm - Piso', 'Piso', 40, '2 CPs - 28 dias'),
], start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=i, column=c, value=v)

C25 = 'FCK 25 | Bombeável - Slump 120±20mm - Brita 1'
C30 = 'FCK 30 | Bombeável - Slump 120±20mm - Brita 0+1'
C40 = 'FCK 40 | Bombeável - Slump 120±20mm - Piso'
ws = wb['Peças']
for i, row in enumerate([
    ('1 a 2', 'Estacas', 'Fundação', 'FCK 25', '455,80', C25, ''),
    ('1 a 2', 'Blocos de Coroamento', 'Fundação', 'FCK 25', '212,50', C25, ''),
    ('1 a 2', 'Térreo (piso)', 'Estrutura', 'FCK 40', '98,32', C40, 'piso polido'),
    ('1 a 2', 'Pavimento Tipo', 'Estrutura', 'FCK 30', '76,40', C30, '12 pavimentos'),
    ('1 a 2', 'Cobertura', 'Estrutura', 'FCK 30', '64,15', C30, ''),
    ('Área Externa', 'Passeio', 'Anexos', 'FCK 25', '35,00', C25, ''),
    ('Área Externa', 'Guarita', 'Anexos', 'FCK 25', '12,80', C25, ''),
    ('Contenção', 'Muro de Arrimo', 'Contenção', 'FCK 30', '58,90', C30, ''),
], start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=i, column=c, value=v)

ws = wb['Usuários']
for i, row in enumerate([
    ('Ana Paula Souza', 'ana.souza@construtoraexemplo.com.br', 'Engenheira Civil', '(11) 98765-4321', 'Engenheiro'),
    ('Carlos Mendes', 'carlos.mendes@construtoraexemplo.com.br', 'Apontador', '(11) 98765-1122', 'Operador de Campo'),
    ('Rita Almeida', 'rita.almeida@construtoraexemplo.com.br', 'Coordenadora da Qualidade', '', 'Qualidade'),
], start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=i, column=c, value=v)

ws = wb['Fornecedores']
for i, row in enumerate([
    ('Usina', 'Concreteira Exemplo S.A.', '98.765.432/0001-10', 'Rod. Anhanguera, km 30', 'João Ribeiro',
     'programacao@concreteiraexemplo.com.br; comercial@concreteiraexemplo.com.br', '(11) 4002-8922', ''),
    ('Laboratório', 'Laboratório Exemplo de Materiais', '11.222.333/0001-44', 'Av. Central, 500', 'Marina Dias',
     'laudos@labexemplo.com.br', '(11) 4002-1234', ''),
], start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=i, column=c, value=v)

wb.save(DST)
print('OK', DST)
