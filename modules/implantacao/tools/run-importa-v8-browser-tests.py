#!/usr/bin/env python3
"""Validação funcional estática do Copilot V8 e do self-test da Pré-análise.

Executa as páginas com os assets locais inline, portanto não depende de servidor HTTP
nem envia planilhas para a rede. Requer Python Playwright, Chromium e openpyxl.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from playwright.sync_api import Page, sync_playwright

ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / "frontend" / "public"
OUT_DIR = ROOT / ".v8-validation"
OUT_DIR.mkdir(exist_ok=True)


def find_example() -> Path:
    candidates = [
        ROOT / "exemplos" / "planilha-cliente-exemplo.xlsx",
        ROOT / "prototype" / "v06" / "exemplos" / "planilha-cliente-exemplo.xlsx",
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError("planilha-cliente-exemplo.xlsx não encontrada")


def inline_page(page_name: str, scripts: list[str]) -> str:
    html = (PUBLIC / page_name).read_text(encoding="utf-8")
    css_path = PUBLIC / "assets" / "design-system.css"
    if css_path.exists():
        css = css_path.read_text(encoding="utf-8")
        css = re.sub(r"@import\s+url\([^;]+;", "", css)
        html = re.sub(
            r'<link[^>]+href=["\']\.?/?assets/design-system\.css["\'][^>]*>',
            f"<style>{css}</style>",
            html,
            flags=re.I,
        )
    for rel in scripts:
        tag = f'<script src="./{rel}"></script>'
        source = (PUBLIC / rel).read_text(encoding="utf-8")
        if tag not in html:
            raise RuntimeError(f"tag não encontrada em {page_name}: {tag}")
        html = html.replace(tag, f"<script>{source}</script>")
    return html


def capture_errors(page: Page) -> list[str]:
    errors: list[str] = []
    page.on("pageerror", lambda exc: errors.append(f"pageerror: {exc}"))

    def on_console(msg: Any) -> None:
        if msg.type == "error" and "Failed to load resource" not in msg.text:
            errors.append(f"console: {msg.text}")

    page.on("console", on_console)
    return errors


def assert_true(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def run_copilot(page: Page, example: Path) -> dict[str, Any]:
    errors = capture_errors(page)
    html = inline_page(
        "importa.html",
        [
            "vendor/xlsx.full.min.js",
            "assets/intake-auditor.js",
            "assets/similarity-engine.js",
        ],
    )
    page.set_content(html, wait_until="load")

    selftest = page.evaluate("() => SelfTest.run()")
    assert_true(selftest["ok"], f"SelfTest do Copilot falhou: {selftest}")
    assert_true(page.locator('[data-action^="mode-"]').count() == 3, "home não exibe os 3 modos")

    page.locator('[data-action="mode-pre"]').click()
    assert_true(page.locator("#filePre").count() == 1, "upload da Pré-importação não abriu")
    page.locator("details summary").click()
    page.locator("#tracosInput").fill(
        "FCK 25 | BRITA 1 | SLUMP 10±2 CM\n"
        "FCK 30 | BRITA 1 | SLUMP 12±2 CM\n"
        "FCK 40 | BRITA 1 | SLUMP 14±2 CM"
    )
    page.locator("#setting-strictMode").select_option("false")
    page.locator("#filePre").set_input_files(str(example))
    page.locator('[data-action="analyze"]').click()
    page.wait_for_function("() => appState.step === 2 && appState.pre.columns.length > 0")

    mapping = page.evaluate(
        "() => ({concreto: appState.pre.mapping.fields.Concreto?.sourceIndex, "
        "data: appState.pre.mapping.fields['Data Concretagem']?.sourceIndex, "
        "nf: appState.pre.mapping.fields['N.F. DANFE']?.sourceIndex})"
    )
    assert_true(mapping == {"concreto": 3, "data": 0, "nf": 1}, f"mapeamento regressivo: {mapping}")

    page.locator('[data-action="prepare-inference"]').click()
    page.wait_for_function("() => appState.step === 3 && appState.pre.candidateRows.length > 0")
    prepared = page.evaluate(
        "() => ({"
        "rows: appState.pre.candidateRows.length,"
        "sourceLines: appState.pre.candidateRows.map(r => r._source.line),"
        "excluded: appState.pre.excludedRows.map(r => ({line:r.line,reason:r.reason})),"
        "proposals: appState.pre.proposals.map(p => ({id:p.id,field:p.field,status:p.status,rule:p.rule}))"
        "})"
    )
    assert_true(prepared["rows"] == 10, f"esperadas 10 concretagens, obtidas {prepared['rows']}")
    assert_true(15 not in prepared["sourceLines"], "linha TOTAL entrou como concretagem")
    assert_true(
        any(x["line"] == 15 and "total" in x["reason"].lower() for x in prepared["excluded"]),
        f"linha TOTAL não foi auditada como excluída: {prepared['excluded']}",
    )

    page.locator('[data-action="accept-all-proposals"]').click()
    pending_fresh = page.evaluate(
        "() => appState.pre.proposals.filter(p => p.status === 'pending' && (p.field === 'Slump' || p.field === 'Flow')).map(p => ({id:p.id,rowId:p.rowId,field:p.field,value:p.value}))"
    )
    assert_true(len(pending_fresh) >= 2, f"faltam propostas de Slump/Flow para revisão: {pending_fresh}")

    accepted = pending_fresh[0]
    rejected = pending_fresh[1]
    page.locator(
        f'[data-action="accept-proposal"][data-proposal="{accepted["id"]}"]'
    ).click()
    page.locator(
        f'[data-action="reject-proposal"][data-proposal="{rejected["id"]}"]'
    ).click()

    decisions = page.evaluate(
        "ids => ids.map(id => { const p=appState.pre.proposals.find(x=>x.id===id); const r=appState.pre.candidateRows.find(x=>x._rowId===p.rowId); return {id,status:p.status,field:p.field,value:r[p.field],provenance:r._provenance[p.field]}; })",
        [accepted["id"], rejected["id"]],
    )
    assert_true(decisions[0]["status"] == "accepted" and decisions[0]["value"] != "", f"aceite não aplicado: {decisions[0]}")
    assert_true(decisions[1]["status"] == "rejected" and decisions[1]["value"] == "", f"rejeição não preservou vazio: {decisions[1]}")

    page.locator('[data-action="finalize-conversion"]').click()
    page.wait_for_function("() => appState.step === 4 && appState.pre.convertedRows.length > 0")
    summary = page.evaluate("() => appState.pre.summary")
    assert_true(summary["total"] == 10, f"resumo final deveria ter 10 linhas: {summary}")

    page.locator('[data-action="go-export"]').click()
    page.wait_for_function("() => appState.step === 5")
    with page.expect_download() as download_info:
        page.locator('[data-action="download-import"]').click()
    download = download_info.value
    export_path = OUT_DIR / download.suggested_filename
    download.save_as(str(export_path))

    wb = load_workbook(export_path, data_only=False, read_only=False)
    assert_true(wb.sheetnames == ["Importação", "Proveniência"], f"abas exportadas incorretas: {wb.sheetnames}")
    ws = wb["Importação"]
    wp = wb["Proveniência"]
    assert_true(ws.max_row == 11, f"exportação deveria conter cabeçalho + 10 linhas: {ws.max_row}")
    assert_true(wp.max_row > 100, f"aba Proveniência inesperadamente curta: {wp.max_row}")
    provenance_headers = [c.value for c in wp[1]]
    assert_true("Proveniência" in provenance_headers and "Aceito pelo Operador" in provenance_headers, f"cabeçalhos de proveniência incompletos: {provenance_headers}")

    page.locator('[data-action="go-home"]').click()
    page.locator('[data-action="mode-post"]').click()
    assert_true(page.locator("#filePost").count() == 1, "Pós-importação não abriu após o fluxo")
    assert_true(not errors, f"erros no navegador: {errors}")

    return {
        "selftest": {"passed": len(selftest["tests"]), "failed": 0},
        "mapping": mapping,
        "candidateRows": prepared["rows"],
        "excludedTotal": next(x for x in prepared["excluded"] if x["line"] == 15),
        "acceptedProposal": decisions[0],
        "rejectedProposal": decisions[1],
        "summary": summary,
        "export": str(export_path),
        "sheets": wb.sheetnames,
        "importRows": ws.max_row - 1,
        "provenanceRows": wp.max_row - 1,
        "browserErrors": errors,
        "postMode": "ok",
    }


def run_intake(page: Page) -> dict[str, Any]:
    errors = capture_errors(page)
    html = inline_page(
        "pre-analise.html",
        [
            "vendor/xlsx.full.min.js",
            "assets/intake-auditor.js",
            "assets/intake-selftest.js",
            "assets/pre-analise.js",
        ],
    )
    page.set_content(html, wait_until="load")
    result = page.evaluate("() => IntakeSelfTest.run(IntakeAuditor, XLSX)")
    assert_true(result["total"] == 27 and result["failed"] == 0, f"Pré-análise regrediu: {result}")
    assert_true(not errors, f"erros na Pré-análise: {errors}")
    return {"passed": result["passed"], "failed": result["failed"], "total": result["total"], "browserErrors": errors}


def main() -> int:
    example = find_example()
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(
            executable_path="/usr/bin/chromium",
            headless=True,
            args=["--no-sandbox", "--disable-dev-shm-usage"],
        )
        context = browser.new_context(accept_downloads=True, viewport={"width": 1440, "height": 1000})
        try:
            copilot = run_copilot(context.new_page(), example)
            intake = run_intake(context.new_page())
        finally:
            context.close()
            browser.close()

    result = {"ok": True, "example": str(example), "copilot": copilot, "intake": intake}
    output = OUT_DIR / "importa-v8-browser-result.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # gate: erro visível e código != 0
        print(f"FALHA: {exc}", file=sys.stderr)
        raise
